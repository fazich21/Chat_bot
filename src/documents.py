"""Read documents from data/ and turn each one into Markdown-style text.

Supported: .pdf, .docx, .md, .txt
Headings are recovered so the chunker can split by section:
  - PDF:  lines in a larger font than the body text become headings
  - DOCX: paragraphs with "Title"/"Heading N" styles become headings
"""
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}


@dataclass
class Document:
    source: str   # file name, e.g. "My Projects.pdf"
    title: str    # document title
    text: str     # Markdown-style text (# title, ## sections, paragraphs separated by blank lines)


# ── PDF ──────────────────────────────────────────────────────────────────
def _pdf_lines(path: Path) -> list[tuple[str, float, int]]:
    """Return (text, font_size, block_id) for every non-empty line in the PDF."""
    import pymupdf

    lines, block_id = [], 0
    with pymupdf.open(path) as pdf:
        for page in pdf:
            for block in page.get_text("dict")["blocks"]:
                block_id += 1
                for line in block.get("lines", []):
                    text = "".join(span["text"] for span in line["spans"]).strip()
                    if text:
                        size = max(span["size"] for span in line["spans"])
                        lines.append((text, round(size, 1), block_id))
    return lines


def read_pdf(path: Path) -> str:
    lines = _pdf_lines(path)
    if not lines:
        return ""

    # Body size = the font size used for the most characters
    weights = Counter()
    for text, size, _ in lines:
        weights[size] += len(text)
    body = weights.most_common(1)[0][0]
    title_size = max(size for _, size, _ in lines)

    out: list[str] = []
    paragraph: list[str] = []
    last_block = None

    def flush():
        if paragraph:
            out.append(" ".join(paragraph))
            paragraph.clear()

    for text, size, block in lines:
        if size < body * 0.75:                  # page headers/footers ("Page 1 of 3")
            continue
        if size >= body * 1.25:                 # heading
            flush()
            level = "#" if size == title_size else "##"
            out.append(f"{level} {text}")
            last_block = None
            continue
        if block != last_block:                 # new PDF block = new paragraph
            flush()
            last_block = block
        paragraph.append(text)
    flush()
    return "\n\n".join(out)


# ── DOCX ─────────────────────────────────────────────────────────────────
def read_docx(path: Path) -> str:
    import docx

    out = []
    for para in docx.Document(str(path)).paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if style == "title" or style == "heading 1":
            out.append(f"# {text}")
        elif style.startswith("heading"):
            out.append(f"## {text}")
        else:
            out.append(text)
    return "\n\n".join(out)


# ── Plain text / Markdown ────────────────────────────────────────────────
def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


READERS = {".pdf": read_pdf, ".docx": read_docx, ".md": read_text, ".txt": read_text}


def read_document(path: Path) -> Document:
    text = READERS[path.suffix.lower()](path)
    title = path.stem
    for line in text.splitlines():
        if line.startswith("# "):
            title = line[2:].strip()
            break
    return Document(source=path.name, title=title, text=text)


def list_data_files(data_dir: Path) -> list[Path]:
    return sorted(p for p in data_dir.iterdir()
                  if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
                  and not p.name.startswith(("~$", ".")))
